from __future__ import annotations

from pathlib import Path
import zipfile
from typing import Any

from openpyxl import load_workbook

from .excel import WrittenField
from .formatters import normalize_text


class ValidationError(RuntimeError):
    pass


def validate_output_file(
    *,
    output_path: Path,
    written_fields: list[WrittenField],
    cell_map: dict[str, Any],
    expected_photo_count: int,
) -> None:
    _validate_zip(output_path)
    workbook = load_workbook(output_path)
    _validate_written_fields(workbook, written_fields)
    _validate_forbidden_values(workbook, cell_map.get("workbook", {}).get("forbidden_values", []))
    _validate_photos(workbook, cell_map, expected_photo_count)


def _validate_zip(output_path: Path) -> None:
    with zipfile.ZipFile(output_path) as archive:
        bad_file = archive.testzip()
    if bad_file:
        raise ValidationError(f"Arquivo XLSX corrompido no item interno: {bad_file}")


def _validate_written_fields(workbook, written_fields: list[WrittenField]) -> None:
    for field in written_fields:
        sheet = workbook[field.sheet]
        actual = sheet[field.cell].value
        expected = field.value
        if actual is None and expected == "":
            continue
        if actual != expected:
            raise ValidationError(
                f"Campo '{field.field_id}' não conferiu em {field.sheet}!{field.cell}. "
                f"Esperado: {expected!r}; encontrado: {actual!r}."
            )


def _validate_forbidden_values(workbook, forbidden_values: list[str]) -> None:
    normalized_forbidden = [normalize_text(value) for value in forbidden_values if str(value).strip()]
    if not normalized_forbidden:
        return
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                text = normalize_text(str(cell.value))
                for forbidden in normalized_forbidden:
                    if forbidden and forbidden in text:
                        raise ValidationError(
                            f"Conteúdo proibido/remanescente encontrado em {sheet.title}!{cell.coordinate}: {cell.value!r}"
                        )


def _validate_photos(workbook, cell_map: dict[str, Any], expected_photo_count: int) -> None:
    if expected_photo_count == 0:
        return
    config = cell_map.get("photos", {})
    sheet_name = config.get("sheet_name")
    if sheet_name:
        sheet = workbook[sheet_name]
    else:
        sheet = next((ws for ws in workbook.worksheets if "foto" in normalize_text(ws.title)), workbook.worksheets[0])
    start_row = _photo_start_row(sheet, config)
    photo_area_count = 0
    for image in sheet._images:
        try:
            anchor_row = image.anchor._from.row + 1
        except AttributeError:
            continue
        if anchor_row >= start_row:
            photo_area_count += 1
    if photo_area_count < expected_photo_count:
        raise ValidationError(
            f"Quantidade de fotos na aba '{sheet.title}' menor que o esperado: "
            f"{photo_area_count} de {expected_photo_count}."
        )


def _photo_start_row(sheet, config: dict[str, Any]) -> int:
    start_cell = config.get("start_cell")
    if start_cell:
        return sheet[start_cell].row
    return 1
