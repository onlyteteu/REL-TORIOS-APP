from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.drawing.image import Image as XlsxImage
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image as PillowImage

from .formatters import format_decimal_br, normalize_text
from .formatters import format_integer_br
from .models import MissingText, RuralVisitReport
from .parser import render_fish_farming, render_livestock


try:
    import yaml
except ImportError as exc:  # pragma: no cover - exercised in installed environments.
    yaml = None
    YAML_IMPORT_ERROR = exc
else:
    YAML_IMPORT_ERROR = None


@dataclass(frozen=True)
class WrittenField:
    field_id: str
    sheet: str
    cell: str
    value: Any


def load_cell_map(path: Path) -> dict[str, Any]:
    if yaml is None:
        raise RuntimeError("PyYAML não está instalado. Execute: pip install -r requirements.txt") from YAML_IMPORT_ERROR
    with path.open("r", encoding="utf-8") as handle:
        data = yaml.safe_load(handle) or {}
    return data


def generate_report_xlsx(
    *,
    template_path: Path,
    output_path: Path,
    report: RuralVisitReport,
    cell_map: dict[str, Any],
    photo_paths: list[Path] | None = None,
) -> list[WrittenField]:
    workbook = load_workbook(template_path)
    written = fill_report_fields(workbook, report, cell_map)
    insert_photos(workbook, cell_map, photo_paths or [])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return written


def fill_report_fields(workbook, report: RuralVisitReport, cell_map: dict[str, Any]) -> list[WrittenField]:
    config = cell_map.get("report_sheet", {})
    sheet = resolve_sheet(workbook, config)
    clear_ranges(sheet, config.get("clear_ranges", []))
    fields = config.get("fields", {})
    written: list[WrittenField] = []
    targets: list[tuple[str, dict[str, Any], Cell]] = []

    for field_id, field_config in fields.items():
        target = resolve_field_cell(sheet, field_config)
        targets.append((field_id, field_config, target))

    for field_id, field_config, target in targets:
        value = render_field(report, field_id)
        if value == MissingText and "missing" in field_config:
            value = field_config["missing"]
        write_cell_preserving_merge(sheet, target, value)
        written.append(WrittenField(field_id=field_id, sheet=sheet.title, cell=target.coordinate, value=value))

    written.extend(fill_machine_table(sheet, report, config.get("machine_table", {})))
    return written


def render_field(report: RuralVisitReport, field_id: str) -> Any:
    if field_id == "client_name":
        return report.client_name
    if field_id == "municipality_uf":
        return report.municipality_uf
    if field_id == "property_name":
        return report.property_name
    if field_id == "discrimination":
        return render_discrimination(report)
    if field_id == "exploitation_type":
        return report.exploitation_type
    if field_id == "productive_situation":
        return report.productive_situation
    if field_id == "own_area":
        return report.own_area.display if report.own_area else MissingText
    if field_id == "leased_area":
        return report.leased_area.display if report.leased_area else MissingText
    if field_id == "total_area":
        return report.total_area.display if report.total_area else MissingText
    if field_id == "own_area_hectares":
        return _decimal_to_float(report.own_area.hectares) if report.own_area else ""
    if field_id == "leased_area_hectares":
        return _decimal_to_float(report.leased_area.hectares) if report.leased_area else ""
    if field_id == "total_area_hectares":
        return _decimal_to_float(report.total_area_hectares.hectares) if report.total_area_hectares else ""
    if field_id == "pasture_area_hectares":
        return _decimal_to_float(report.pasture_area_hectares.hectares) if report.pasture_area_hectares else ""
    if field_id == "crop_area_hectares":
        return _decimal_to_float(report.crop_area_hectares.hectares) if report.crop_area_hectares else ""
    if field_id == "activities":
        return ", ".join(report.activities) if report.activities else MissingText
    if field_id == "main_activity":
        return report.main_activity
    if field_id == "main_cultures":
        return render_main_cultures(report)
    if field_id == "livestock":
        return render_livestock(report.livestock)
    if field_id == "fish_farming":
        return render_fish_farming(report.fish_farming)
    if field_id == "pastures":
        return ", ".join(report.forage_species) if report.forage_species else MissingText
    if field_id == "improvements":
        return "\n".join(report.improvements) if report.improvements else MissingText
    if field_id == "machines":
        return "\n".join(
            f"Descrição: {item.description}; Fabricante: {item.manufacturer}; Modelo: {item.model}"
            for item in report.machines
        )
    if field_id == "machine_description":
        return report.machines[0].description if report.machines else MissingText
    if field_id == "machine_manufacturer":
        return report.machines[0].manufacturer if report.machines else "-"
    if field_id == "machine_model":
        return report.machines[0].model if report.machines else "-"
    if field_id == "property_characterization":
        return render_property_characterization(report)
    if field_id == "infrastructure_text":
        return render_infrastructure_text(report)
    if field_id == "activity_comments":
        return render_activity_comments(report)
    if field_id == "summary_comments":
        return render_summary_comments(report)
    if field_id == "other_comments":
        return render_other_comments(report)
    if field_id == "final_conclusion":
        return render_final_conclusion(report)
    if field_id == "technical_comments":
        return report.technical_comments
    if field_id == "conclusion":
        return report.conclusion
    return MissingText


def _decimal_to_float(value: Decimal) -> float:
    return float(value)


def render_discrimination(report: RuralVisitReport) -> str:
    lines = [
        f"Nome da propriedade: {report.property_name}",
        f"Tipo de exploração: {report.exploitation_type}",
        f"Atividades desenvolvidas: {', '.join(report.activities) if report.activities else MissingText}",
        f"Situação produtiva: {report.productive_situation}",
    ]
    return "\n".join(lines)


def render_main_cultures(report: RuralVisitReport) -> str:
    if not report.main_cultures:
        return MissingText
    return ", ".join(report.main_cultures)


def render_property_characterization(report: RuralVisitReport) -> str:
    return render_infrastructure_text(report)


def render_infrastructure_text(report: RuralVisitReport) -> str:
    sentences = []
    if report.property_name != MissingText:
        sentences.append(f"Propriedade informada: {report.property_name}.")
    if report.activities:
        sentences.append(f"Tipo de uso informado: {', '.join(report.activities)}.")
    if report.paddocks:
        sentences.append(f"Piquetes informados: {'; '.join(report.paddocks)}.")
    if report.forage_species:
        sentences.append(f"Espécies forrageiras informadas: {', '.join(report.forage_species)}.")
    if report.troughs:
        sentences.append(f"Cochos/bebedouros informados: {'; '.join(report.troughs)}.")
    if report.fences:
        sentences.append(f"Cercas informadas: {'; '.join(report.fences)}.")
    if report.corrals:
        sentences.append(f"Currais informados: {'; '.join(report.corrals)}.")
    if report.pasture_conditions:
        sentences.append(f"Condições das pastagens informadas: {'; '.join(report.pasture_conditions)}.")
    if report.crops:
        sentences.append(f"Culturas existentes informadas: {'; '.join(report.crops)}.")
    if report.support_structures:
        sentences.append(f"Estruturas de apoio informadas: {'; '.join(report.support_structures)}.")
    if report.water_resources:
        water_text = "; ".join(item.rstrip(".") for item in report.water_resources)
        sentences.append(f"Recursos hídricos informados: {water_text}.")
    return " ".join(sentences) if sentences else MissingText


def render_activity_comments(report: RuralVisitReport) -> str:
    sections = []
    if report.livestock:
        sections.append(render_livestock(report.livestock))
    fish = render_fish_farming(report.fish_farming)
    if fish != MissingText:
        sections.append(fish)
    if report.forage_species:
        sections.append(f"Espécies forrageiras informadas: {', '.join(report.forage_species)}.")
    if report.herd_movements:
        sections.append(f"Movimentação do rebanho informada: {'; '.join(report.herd_movements)}.")
    if not sections:
        return MissingText
    return " ".join(sections) + " Informações registradas conforme anotações fornecidas."


def render_summary_comments(report: RuralVisitReport) -> str:
    if not report.activities:
        return MissingText
    parts = [f"PROPRIEDADE COM ATIVIDADE INFORMADA DE {', '.join(report.activities).upper()}"]
    if report.total_area_hectares:
        parts.append(f"EM ÁREA TOTAL INFORMADA DE {format_decimal_br(report.total_area_hectares.hectares).upper()} HA")
    if report.livestock:
        livestock_counts = [item.count_heads for item in report.livestock if item.count_heads is not None]
        if livestock_counts:
            parts.append(f"COM REBANHO INFORMADO DE {format_integer_br(sum(livestock_counts))} CABEÇAS")
    return ". ".join(parts) + "."


def render_other_comments(report: RuralVisitReport) -> str:
    comments = []
    if report.activities:
        comments.append(f"A exploração rural informada compreende {', '.join(report.activities)}.")
    if report.total_area_hectares:
        comments.append(f"A área total registrada nas anotações corresponde a {format_decimal_br(report.total_area_hectares.hectares)} ha.")
    if report.herd_movements:
        comments.append(f"Movimentação do rebanho informada: {'; '.join(report.herd_movements)}.")
    if report.machines:
        comments.append(render_field(report, "machines"))
    return " ".join(comments) if comments else MissingText


def render_final_conclusion(report: RuralVisitReport) -> str:
    if not report.activities and not report.total_area_hectares:
        return MissingText
    return (
        "Conclui-se que o presente registro foi elaborado com base exclusivamente nas informações "
        "fornecidas para a visita rural, sem inclusão de dados não informados pelo usuário."
    )


def resolve_sheet(workbook, config: dict[str, Any]) -> Worksheet:
    sheet_name = config.get("sheet_name")
    if sheet_name:
        return workbook[sheet_name]
    contains = normalize_text(config.get("sheet_name_contains", ""))
    if contains:
        for sheet in workbook.worksheets:
            if contains in normalize_text(sheet.title):
                return sheet
    return workbook.worksheets[0]


def resolve_photo_sheet(workbook, config: dict[str, Any]) -> Worksheet:
    sheet_name = config.get("sheet_name")
    if sheet_name:
        return workbook[sheet_name]
    for sheet in workbook.worksheets:
        if "foto" in normalize_text(sheet.title):
            return sheet
    return workbook.worksheets[0]


def resolve_field_cell(sheet: Worksheet, field_config: dict[str, Any]) -> Cell:
    if field_config.get("cell"):
        return writable_cell(sheet, sheet[field_config["cell"]])

    label = field_config.get("label")
    if not label:
        raise ValueError("Campo sem 'cell' ou 'label' no cell_map.yaml.")
    label_cell = find_label_cell(sheet, label)
    if label_cell is None:
        raise ValueError(f"Rótulo não encontrado na aba '{sheet.title}': {label}")
    row_offset, col_offset = field_config.get("offset", [0, 1])
    return writable_cell(sheet, sheet.cell(label_cell.row + row_offset, label_cell.column + col_offset))


def find_label_cell(sheet: Worksheet, label: str) -> Cell | None:
    target = normalize_text(label)
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            text = normalize_text(str(cell.value))
            if target == text or target in text:
                return writable_cell(sheet, cell)
    return None


def writable_cell(sheet: Worksheet, cell: Cell) -> Cell:
    for merged in sheet.merged_cells.ranges:
        if cell.coordinate in merged:
            return sheet.cell(merged.min_row, merged.min_col)
    return cell


def write_cell_preserving_merge(sheet: Worksheet, cell: Cell, value: Any) -> None:
    target = writable_cell(sheet, cell)
    target.value = value
    target.alignment = target.alignment.copy(wrap_text=True, vertical="top")


def clear_ranges(sheet: Worksheet, ranges: list[str]) -> None:
    for range_address in ranges:
        for row in sheet[range_address]:
            for cell in row:
                if cell.__class__.__name__ == "MergedCell":
                    continue
                cell.value = None


def insert_photos(workbook, cell_map: dict[str, Any], photo_paths: list[Path]) -> None:
    if not photo_paths:
        return

    config = cell_map.get("photos", {})
    sheet = resolve_photo_sheet(workbook, config)
    start_cell = _resolve_photo_start_cell(sheet, config)

    if config.get("clear_existing_photos", True):
        _clear_existing_photos(sheet, start_cell.row)

    columns = int(config.get("columns", 1))
    row_step = int(config.get("row_step", 19))
    col_step = int(config.get("col_step", 0))
    max_width = int(config.get("image_max_width_px", 640))
    max_height = int(config.get("image_max_height_px", 360))
    caption_row_offset = int(config.get("caption_row_offset", 18))
    caption_col_offset = int(config.get("caption_col_offset", 6))

    for index, photo_path in enumerate(photo_paths):
        if not photo_path.exists():
            raise FileNotFoundError(f"Foto não encontrada: {photo_path}")
        row = start_cell.row + (index // columns) * row_step
        col = start_cell.column + (index % columns) * col_step
        width, height = fit_image_size(photo_path, max_width, max_height)

        image = XlsxImage(str(photo_path))
        image.width = width
        image.height = height
        image.anchor = f"{get_column_letter(col)}{row}"
        sheet.add_image(image)

        caption_cell = writable_cell(sheet, sheet.cell(row + caption_row_offset, col + caption_col_offset))
        caption_cell.value = caption_for_photo(photo_path, index + 1)
        caption_cell.alignment = caption_cell.alignment.copy(horizontal="center", vertical="center", wrap_text=True)


def _resolve_photo_start_cell(sheet: Worksheet, config: dict[str, Any]) -> Cell:
    if config.get("start_cell"):
        return writable_cell(sheet, sheet[config["start_cell"]])
    label = config.get("section_label")
    if label:
        label_cell = find_label_cell(sheet, label)
        if label_cell:
            row_offset, col_offset = config.get("offset_from_label", [2, 0])
            return writable_cell(sheet, sheet.cell(label_cell.row + row_offset, label_cell.column + col_offset))
    return writable_cell(sheet, sheet["D209"])


def _clear_existing_photos(sheet: Worksheet, start_row: int) -> None:
    kept_images = []
    for image in sheet._images:
        try:
            anchor_row = image.anchor._from.row + 1
        except AttributeError:
            kept_images.append(image)
            continue
        if anchor_row < start_row:
            kept_images.append(image)
    sheet._images = kept_images


def fit_image_size(path: Path, max_width: int, max_height: int) -> tuple[int, int]:
    with PillowImage.open(path) as image:
        original_width, original_height = image.size
    scale = min(max_width / original_width, max_height / original_height)
    return max(1, int(original_width * scale)), max(1, int(original_height * scale))


def caption_for_photo(path: Path, index: int) -> str:
    stem = path.stem.strip().replace("_", " ").replace("-", " ")
    if stem:
        return stem[:80]
    return f"Registro fotográfico {index}"


def fill_machine_table(sheet: Worksheet, report: RuralVisitReport, config: dict[str, Any]) -> list[WrittenField]:
    if not config:
        return []
    start_row = int(config.get("start_row", 40))
    max_rows = int(config.get("max_rows", 15))
    columns = config.get("columns", {"description": "A", "manufacturer": "E", "model": "F"})
    written = []
    machines = report.machines[:max_rows] if report.machines else []
    if not machines:
        machines = []
    for offset, machine in enumerate(machines):
        row = start_row + offset
        values = {
            "description": machine.description or MissingText,
            "manufacturer": machine.manufacturer or "-",
            "model": machine.model or "-",
        }
        for key, value in values.items():
            coord = f"{columns[key]}{row}"
            cell = writable_cell(sheet, sheet[coord])
            write_cell_preserving_merge(sheet, cell, value)
            written.append(WrittenField(field_id=f"machine_{key}_{offset + 1}", sheet=sheet.title, cell=cell.coordinate, value=value))
    return written


def safe_output_name(pattern: str, client_name: str) -> str:
    safe_client = re.sub(r'[<>:"/\\|?*]+', "", client_name).strip()
    safe_client = re.sub(r"\s+", " ", safe_client)
    filename = pattern.format(client_name=safe_client or "Não informado")
    filename = re.sub(r'[<>:"/\\|?*]+', "", filename).strip()
    return filename or "RELATÓRIO DE VISITA - Não informado.xlsx"
