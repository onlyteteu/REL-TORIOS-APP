from __future__ import annotations

from pathlib import Path
import json

from openpyxl import load_workbook


def inspect_workbook_labels(workbook_path: Path) -> dict:
    workbook = load_workbook(workbook_path, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        labels = []
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip():
                    labels.append({"cell": cell.coordinate, "value": cell.value.strip()})
        sheets.append(
            {
                "title": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
                "labels": labels,
            }
        )
    return {"sheets": sheets}


def write_inspection_json(workbook_path: Path, output_path: Path) -> None:
    data = inspect_workbook_labels(workbook_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
